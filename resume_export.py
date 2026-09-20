"""Create safe comparisons and clean downloadable resume documents."""

from difflib import SequenceMatcher
from html import escape
from io import BytesIO

from docx import Document
from docx.shared import Pt


def build_comparison_html(original: str, updated: str) -> str:
    """Return an escaped side-by-side line comparison with changed rows highlighted."""

    old_lines = original.splitlines()
    new_lines = updated.splitlines()
    matcher = SequenceMatcher(None, old_lines, new_lines)
    rows: list[str] = []

    def add_row(old: str, new: str, old_class: str = "", new_class: str = ""):
        rows.append(
            "<tr>"
            f'<td class="{old_class}">{escape(old) or "&nbsp;"}</td>'
            f'<td class="{new_class}">{escape(new) or "&nbsp;"}</td>'
            "</tr>"
        )

    for tag, old_start, old_end, new_start, new_end in matcher.get_opcodes():
        old_chunk = old_lines[old_start:old_end]
        new_chunk = new_lines[new_start:new_end]
        row_count = max(len(old_chunk), len(new_chunk))
        for index in range(row_count):
            old_line = old_chunk[index] if index < len(old_chunk) else ""
            new_line = new_chunk[index] if index < len(new_chunk) else ""
            if tag == "equal":
                add_row(old_line, new_line)
            elif tag == "delete":
                add_row(old_line, "", "removed", "")
            elif tag == "insert":
                add_row("", new_line, "", "added")
            else:
                add_row(old_line, new_line, "removed", "added")

    return """
    <style>
      .resume-diff {width:100%; border-collapse:collapse; table-layout:fixed;}
      .resume-diff th {padding:10px; text-align:left; background:#f0f2f6;}
      .resume-diff td {width:50%; padding:7px 10px; border:1px solid #ddd;
                       vertical-align:top; white-space:pre-wrap; overflow-wrap:anywhere;}
      .resume-diff .removed {background:#ffe3e3; color:#8b0000;}
      .resume-diff .added {background:#dcfce7; color:#14532d;}
    </style>
    <table class="resume-diff">
      <thead><tr><th>Original resume</th><th>Updated resume</th></tr></thead>
      <tbody>""" + "".join(rows) + "</tbody></table>"


def create_resume_docx(resume_text: str) -> bytes:
    """Create a clean DOCX containing no comparison colors or highlights."""

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    for raw_line in resume_text.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
        elif line.startswith(("- ", "• ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif len(line) <= 45 and (line.isupper() or line.endswith(":")):
            document.add_heading(line.rstrip(":"), level=1)
        else:
            document.add_paragraph(line)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
